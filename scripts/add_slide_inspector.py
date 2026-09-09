# -*- coding: utf-8 -*-
"""
기존 13장 발표 덱(발표_Career-Path-Navigator.pptx)을 *그대로 두고*,
Document Inspector(두 문서 유사도 분석) 슬라이드 1장을 12번(라이브 데모) 뒤에 끼운
새 파일을 생성한다. 디자인 상수/헬퍼는 make_ppt.py와 동일(테마 일치).

원본 미수정 — 출력은 별도 파일.
실행: python scripts/add_slide_inspector.py
출력: 발표_Career-Path-Navigator_유사도추가.pptx (프로젝트 루트)
"""
import os
import copy
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---- 색상 (make_ppt.py와 동일) ----
INK    = RGBColor(0x0F, 0x17, 0x2A)
NAVY   = RGBColor(0x1E, 0x3A, 0x8A)
BLUE   = RGBColor(0x25, 0x63, 0xEB)
SKY    = RGBColor(0x60, 0xA5, 0xFA)
PALE   = RGBColor(0xDB, 0xEA, 0xFE)
AMBER  = RGBColor(0xB4, 0x53, 0x09)
GREY   = RGBColor(0x47, 0x55, 0x69)
FAINT  = RGBColor(0x94, 0xA3, 0xB8)
HAIR   = RGBColor(0xE2, 0xE8, 0xF0)
ZEBRA  = RGBColor(0xF8, 0xFA, 0xFC)
CARD   = RGBColor(0xF1, 0xF5, 0xF9)
HLROW  = RGBColor(0xFE, 0xF3, 0xC7)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "맑은 고딕"

M  = 0.65
CW = 13.333 - 2 * M

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC  = os.path.join(ROOT, "발표_Career-Path-Navigator.pptx")
OUT  = os.path.join(ROOT, "발표_Career-Path-Navigator_유사도추가.pptx")


# ---- 헬퍼 (make_ppt.py와 동일) ----
def _ea(run):
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", FONT)


def style_run(run, size, bold=False, color=INK):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    _ea(run)


def add_box(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0
    tf.margin_top = 0; tf.margin_bottom = 0
    return tf


def add_shape(slide, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE,
              line_color=None, line_w=0.75):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line_color is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line_color; sp.line.width = Pt(line_w)
    sp.shadow.inherit = False
    return sp


def card(slide, x, y, w, h, fill=CARD, line=None, line_w=0.75, radius=0.07):
    sp = add_shape(slide, x, y, w, h, fill, MSO_SHAPE.ROUNDED_RECTANGLE,
                   line_color=line, line_w=line_w)
    try:
        sp.adjustments[0] = radius
    except Exception:
        pass
    tf = sp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.16); tf.margin_right = Inches(0.16)
    tf.margin_top = Inches(0.06); tf.margin_bottom = Inches(0.06)
    return sp, tf


def para(tf, runs, align=PP_ALIGN.LEFT, before=0, after=0, line=1.04, first=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    if before: p.space_before = Pt(before)
    if after:  p.space_after = Pt(after)
    p.line_spacing = line
    for (txt, size, bold, color) in runs:
        r = p.add_run(); style_run(r, size, bold, color); r.text = txt
    return p


def notes(slide, text, template_slide=None):
    ns = slide.notes_slide
    tf = ns.notes_text_frame
    if tf is None and template_slide is not None:
        # 노트 placeholder가 없으면 노트 있는 기존 슬라이드에서 body placeholder 복제
        src_ph = template_slide.notes_slide.notes_placeholder
        if src_ph is not None:
            ns.shapes._spTree.append(copy.deepcopy(src_ph._element))
            tf = ns.notes_text_frame
    if tf is not None:
        tf.text = text
    else:
        print("  [warn] notes placeholder 없음 — 노트 생략(발표대본.md 참고)")


def header(slide, kicker, title, star=False):
    tf = add_box(slide, M, 0.46, CW, 0.34)
    para(tf, [(kicker, 13, True, BLUE)], first=True)
    tf2 = add_box(slide, M, 0.80, CW, 0.72)
    runs = [(title, 30, True, INK)]
    if star:
        runs.append(("  ★", 21, True, AMBER))
    para(tf2, runs, first=True)
    add_shape(slide, M, 1.60, 1.05, 0.045, BLUE)


def footer_label(slide):
    # 원본 13장의 페이지번호 체계를 흩뜨리지 않도록, 새 슬라이드는 좌측 라벨만.
    tfL = add_box(slide, M, 7.08, 7, 0.3)
    para(tfL, [("Career-Path Navigator", 9, False, FAINT)], first=True)


# ============================================================
# 새 슬라이드 — Document Inspector (두 문서 유사도 분석)
# ============================================================
prs = Presentation(SRC)
BLANK = prs.slide_layouts[6]
s = prs.slides.add_slide(BLANK)

header(s, "데모", "Document Inspector — 두 자소서 유사도 분석")

# 좌측: 분석 4요소
lw = 6.05
items = [
    ("전체 문서 유사도", "Bag-of-N-grams 코사인 유사도 (0~100%)"),
    ("공통 키워드 · n-gram", "어미·조사 제외 → 도메인 핵심 어휘 자동 추출"),
    ("문장 단위 매칭", "cosine ≥ 0.25 인 문장 쌍을 좌우 대조"),
    ("도메인 분포 비교", "두 문서의 Hybrid score 막대그래프 병치"),
]
for i, (t, d) in enumerate(items):
    cy = 1.98 + i * 1.16
    sp, tfc = card(s, M, cy, lw, 0.98, fill=ZEBRA, line=HAIR)
    tfc.margin_left = Inches(0.78)
    para(tfc, [(t, 16, True, INK)], first=True)
    para(tfc, [(d, 12.5, False, GREY)], before=3)
    nsp = add_shape(s, M + 0.18, cy + 0.26, 0.46, 0.46, BLUE, shape=MSO_SHAPE.OVAL)
    ntf = nsp.text_frame; ntf.word_wrap = False
    ntf.margin_left = 0; ntf.margin_right = 0; ntf.margin_top = 0; ntf.margin_bottom = 0
    ntf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(ntf, [(str(i + 1), 16, True, WHITE)], align=PP_ALIGN.CENTER, first=True)

# 우측: 실제 예시 카드
rx = M + lw + 0.45
rw = 12.033 - lw - 0.45
sp, tfc = card(s, rx, 1.98, rw, 3.66, fill=WHITE, line=FAINT, line_w=1.0)
tfc.vertical_anchor = MSO_ANCHOR.TOP
tfc.margin_left = Inches(0.26); tfc.margin_right = Inches(0.26)
tfc.margin_top = Inches(0.20)
para(tfc, [("예시 — 같은 간호 직군, 다른 진로", 13, True, BLUE)], first=True)
para(tfc, [("연세대 간호대학원  ↔  중앙대병원 간호직", 13.5, True, INK)], before=4)
para(tfc, [("(진학 자소서)            (취업 자소서)", 11, False, FAINT)], before=2)
para(tfc, [("45%", 46, True, BLUE)], align=PP_ALIGN.CENTER, before=10)
para(tfc, [("BoN 코사인 전체 유사도", 12, False, GREY)],
     align=PP_ALIGN.CENTER, before=0)
para(tfc, [("공통 키워드", 12, True, NAVY)], before=12)
para(tfc, [("간호 · 간호사 · 전문 · 환자", 14, True, INK)], before=2)
para(tfc, [("도메인 분포 — 두 문서 모두 Nursing 최상위", 12, False, GREY)], before=8)

# 하단 강조 바
sp, tfc = card(s, M, 5.92, 12.033, 0.72, fill=HLROW)
para(tfc, [("분류를 넘어 ", 14.5, True, AMBER),
           ("문서 간 관계까지 정량화", 14.5, True, AMBER),
           (" — 표절·중복·유사 지원자 분석으로 확장 가능", 14.5, False, INK)],
     align=PP_ALIGN.CENTER, first=True)

footer_label(s)
notes(s, "[진행 방법] Document Inspector 탭으로 이동 → 왼쪽 '문서 A'에 연세대 간호대학원 "
         "자소서, 오른쪽 '문서 B'에 중앙대병원 간호직 자소서를 각각 올리고 → 가운데 '유사도 "
         "분석' 버튼을 누르면 결과가 바로 뜹니다.\n\n"
         "[설명 멘트] 분류만이 아니라 문서 사이의 관계도 분석합니다. 방금처럼 자소서 두 개를 "
         "양쪽에 올리고 분석을 누르면 되는데요, 지금 올린 두 문서는 둘 다 간호 직군이지만 하나는 "
         "대학원 진학, 하나는 병원 취업 자소서입니다. 결과를 보면 BoN 코사인 전체 유사도가 45%로 "
         "나오고, 그 아래 공통 키워드로 간호·간호사·환자 같은 도메인 어휘가 자동 추출됩니다. "
         "조사나 어미가 아니라 의미 있는 전문 어휘가 잡힙니다. 맨 아래 도메인 분포 막대그래프에서는 "
         "두 문서 모두 Nursing에 쏠려, 다른 기관·다른 진로지만 같은 직군임이 한눈에 보입니다. "
         "이렇게 문서 비교는 표절·중복 검사나 유사 지원자 분석에도 응용할 수 있습니다.",
      template_slide=prs.slides[11])

# ---- 새 슬라이드를 12번(라이브 데모) 뒤, 마지막 마무리 슬라이드 앞으로 이동 ----
# 원본 13장: ... index11=Dashboard(12번), index12=마무리(13번). 새 슬라이드는 맨 끝(index13).
xml_slides = prs.slides._sldIdLst
slide_ids = list(xml_slides)
new_id = slide_ids[-1]
xml_slides.remove(new_id)
xml_slides.insert(12, new_id)   # Dashboard 바로 뒤

prs.save(OUT)
print("SAVED:", OUT, "| slides:", len(list(prs.slides)))

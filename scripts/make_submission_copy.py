# -*- coding: utf-8 -*-
"""
발표_Career-Path-Navigator_유사도추가.pptx 와 내용은 동일하되,
모든 발표자 노트(대본)를 비운 *제출용* 버전을 새 파일로 만든다.
원본/유사도추가 파일은 미수정.
실행: python scripts/make_submission_copy.py
출력: 발표_Career-Path-Navigator_유사도추가_제출용.pptx
"""
import os
from pptx import Presentation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC  = os.path.join(ROOT, "발표_Career-Path-Navigator_유사도추가.pptx")
OUT  = os.path.join(ROOT, "발표_Career-Path-Navigator_유사도추가_제출용.pptx")

prs = Presentation(SRC)
cleared = 0
for s in prs.slides:
    if s.has_notes_slide:
        tf = s.notes_slide.notes_text_frame
        if tf is not None and tf.text.strip():
            tf.text = ""
            cleared += 1

prs.save(OUT)
print("SAVED:", OUT)
print("slides:", len(list(prs.slides)), "| 노트 비운 슬라이드 수:", cleared)
